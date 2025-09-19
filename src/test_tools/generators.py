'''
 * This file is part of Fika.
 *
 * Fika is free software; you can redistribute it and/or modify
 * it under the terms of the GNU General Public License as published by
 * the Free Software Foundation; version 2 of the License.
 *
 * Fika is distributed in the hope that it will be useful,
 * but WITHOUT ANY WARRANTY; without even the implied warranty of
 * MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
 * GNU General Public License for more details.
 *
 * You should have received a copy of the GNU General Public License
 * along with Fika.  If not, see <http://www.gnu.org/licenses/>.
'''

import os
from PIL import Image
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
import zipfile
import wave

OUTDIR = os.path.join(os.path.dirname(__file__), "synthetic_files")
os.makedirs(OUTDIR, exist_ok=True)

def generate_png(path=f"{OUTDIR}/test.png"):
    Image.new("RGB", (1, 1), color="red").save(path)

def generate_jpg(path=f"{OUTDIR}/test.jpg"):
    Image.new("RGB", (1, 1), color="blue").save(path, "JPEG")

def generate_gif(path=f"{OUTDIR}/test.gif"):
    Image.new("RGB", (1, 1), color="green").save(path, "GIF")

def generate_bmp(path=f"{OUTDIR}/test.bmp"):
    Image.new("RGB", (1, 1), color="yellow").save(path, "BMP")

def generate_pdf(path=f"{OUTDIR}/test.pdf"):
    c = canvas.Canvas(path)
    c.drawString(100, 750, "Hello PDF")
    c.save()

def generate_docx(path=f"{OUTDIR}/test.docx"):
    doc = Document()
    doc.add_paragraph("Hello DOCX")
    doc.save(path)

def generate_xlsx(path=f"{OUTDIR}/test.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Hello XLSX"
    wb.save(path)

def generate_zip(path=f"{OUTDIR}/test.zip"):
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("hello.txt", "Hello ZIP")

def generate_wav(path=f"{OUTDIR}/test.wav"):
    with wave.open(path, "w") as wf:
        wf.setnchannels(1)
        wf.setsampwidth(2)
        wf.setframerate(44100)
        silence = (b"\x00\x00" * 44100)
        wf.writeframes(silence)

# --- Optional generators (skip if deps missing) ---
def generate_mp4(path=f"{OUTDIR}/test.mp4"):
    try:
        import moviepy.editor as mpy
    except ImportError:
        print("⚠️ Skipping MP4 (moviepy not installed)")
        return
    clip = mpy.ColorClip(size=(640, 480), color=(255, 0, 0), duration=2)
    clip.write_videofile(path, fps=24, codec="libx264", logger=None)

def generate_pptx(path=f"{OUTDIR}/test.pptx"):
    try:
        from pptx import Presentation
    except ImportError:
        print("⚠️ Skipping PPTX (python-pptx not installed)")
        return
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Hello PPTX"
    prs.save(path)

def generate_all():
    generate_png()
    generate_jpg()
    generate_gif()
    generate_bmp()
    generate_pdf()
    generate_docx()
    generate_xlsx()
    generate_zip()
    generate_wav()
    generate_mp4()
    generate_pptx()
    print(f"Synthetic files generated in {OUTDIR}")
